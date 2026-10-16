from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


REPO = Path(__file__).resolve().parents[1]
SCRIPT = REPO / "plugins/skill-orchestrator/skills/seedance-shot-prompt-compiler/scripts/seedance_shot_prompt_compiler.py"
SPEC = importlib.util.spec_from_file_location("seedance_prompt_compiler_test", SCRIPT)
compiler = importlib.util.module_from_spec(SPEC); assert SPEC.loader; SPEC.loader.exec_module(compiler)


def package() -> dict:
    shot = {
        "shot_id": "E001S001C001", "framing": "中景", "movement": "固定", "angle": "平视",
        "camera_position": "门外看向门内", "lens_tendency": "50mm定焦镜头（全画幅等效）",
        "blocking": "甲在左，乙在右", "axis": "门床轴线", "eyeline": "甲→乙",
        "reaction": "乙停下", "performance_target": "甲保持克制", "dramatic_purpose": "建立关系",
        "first_state": "甲站定", "last_state": "乙停下", "subject_action": "中景。前景是门框；中景甲站在左侧，乙从右侧走入；后景是房门。",
        "dialogue": "甲：‘这是完整而且很长的一句台词。’", "sound_effects": "房内环境声。",
    }
    value = {
        "schema_version": "SceneDirectionPackageV1", "status": "sealed", "version": "v1",
        "integrity": {"canonical_json_sha256": None},
        "scenes": [{"scene_id": "scn-001", "production_scene_code": "E001S001", "worksheet_name": "S001_测试", "shots": [shot]}],
    }
    value["integrity"]["canonical_json_sha256"] = compiler.package_digest(value)
    return value


def records() -> dict:
    return {
        "schema_version": "SeedanceShotPromptRecordCollectionV1",
        "records": [{
            "shot_id": "E001S001C001",
            "visual": "中景。前景是门框，中景甲站在左侧，乙从右侧走入，后景是房门。画面结束时乙停下。",
            "prompt": "甲站在画面左侧，乙从右侧走入。50mm固定镜头从门外平视门内；甲控制呼吸，目光停在乙脸上，说：“这是完整而且很长的一句台词。”房内环境声持续，乙随后停住脚步。最后保持两人距离和门床轴线。",
        }],
    }


def v2_records() -> dict:
    return {
        "schema_version": "SeedanceShotPromptRecordCollectionV2",
        "review_label": "正式版",
        "records": [{
            "shot_id": "E001S001C001",
            "visual": "中景。前景是门框，中景甲站在左侧，乙从右侧走入，后景是房门。画面结束时乙停下。",
            "performance_beats": [{
                "beat_id": "B01", "trigger": "乙进入后", "dialogue_fragment": "这是完整而且很长的一句台词。",
                "visible_action": "甲控制呼吸并看向乙", "micro_expression": "目光停在乙脸上",
                "voice": "语速平稳", "end_state": "乙停下", "camera_segment_id": "A",
            }],
            "sound_design": {
                "ambient_bed": "房内环境声。", "spot_effects": ["乙停步的鞋底轻响"],
                "silence_beats": ["台词后留短暂停顿"],
            },
            "camera_segments": [{
                "segment_id": "A", "framing": "中景", "movement": "固定机位", "angle": "平视",
                "camera_position": "门外看向门内", "lens": "50mm定焦镜头（全画幅等效）", "axis": "门床轴线",
            }],
            "prompt": "甲站在画面左侧，乙从右侧走入。摄影机以50mm定焦镜头（全画幅等效）从门外平视门内，固定机位。房内环境声。甲控制呼吸并看向乙，目光停在乙脸上，以平稳语速说：‘这是完整而且很长的一句台词。’乙停下（乙停步的鞋底轻响）；台词后留短暂停顿。",
        }],
    }


def reference(path: Path, title: str = "镜头版式模板") -> None:
    wb = Workbook(); ws = wb.active; ws.title = title
    ws.merge_cells("A1:K1"); ws.merge_cells("N1:R1")
    for row in range(1, 6):
        ws.row_dimensions[row].height = 24 if row < 3 else 120
        for col in range(1, 19):
            if not (row == 1 and (2 <= col <= 11 or 15 <= col <= 18)):
                ws.cell(row, col, "模板" if row < 3 else "")
    for name, columns in (("决定日志", 8), ("来源追踪", 5)):
        audit = wb.create_sheet(name)
        audit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        for row in range(1, 4):
            for col in range(1, columns + 1):
                if row != 1 or col == 1:
                    audit.cell(row, col, "值")
    wb.save(path)


class SeedancePromptCompilerTests(unittest.TestCase):
    def test_v2_schema_is_packaged_and_names_required_structures(self) -> None:
        schema_path = SCRIPT.parents[1] / "schemas/shot-prompt-records-v2.schema.json"
        schema = json.loads(schema_path.read_text(encoding="utf-8"))
        self.assertEqual(schema["properties"]["schema_version"]["const"], "SeedanceShotPromptRecordCollectionV2")
        self.assertIn("performance_beats", schema["$defs"]["record"]["required"])
        self.assertIn("sound_design", schema["$defs"]["record"]["required"])
        self.assertIn("camera_segments", schema["$defs"]["record"]["required"])

    def test_rejects_unsealed_package(self) -> None:
        value = package(); value["status"] = "draft"
        self.assertTrue(any("sealed" in error for error in compiler.validate_package(value)))

    def test_rejects_timecodes_and_image_references(self) -> None:
        value = records(); value["records"][0]["prompt"] = "@图片1，0秒-3秒，甲：‘这是完整而且很长的一句台词。’50mm镜头。"
        errors = compiler.validate_records(package(), value)
        self.assertTrue(any("image-reference" in error for error in errors))
        self.assertTrue(any("timecode" in error for error in errors))

    def test_requires_complete_dialogue_and_sound_inside_prompt(self) -> None:
        value = records(); value["records"][0]["prompt"] = "甲走入，50mm固定镜头平视拍摄。"
        errors = compiler.validate_records(package(), value)
        self.assertTrue(any("complete dialogue" in error for error in errors))
        self.assertTrue(any("sound clause" in error for error in errors))

    def test_rejects_generic_prefix_and_audit_tail(self) -> None:
        value = records()
        value["records"][0]["prompt"] = "16:9古装写实电影画面。" + value["records"][0]["prompt"] + "台词口型与台词列同步，不新增对白。"
        errors = compiler.validate_records(package(), value)
        self.assertTrue(any("audit-style" in error for error in errors))

    def test_renders_chinese_review_workbook_without_duration_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root / "reference.xlsx"; output = root / "review.xlsx"
            reference(source)
            compiler.render_workbook(package(), records(), source, output)
            wb = load_workbook(output)
            ws = wb["S001_测试"]
            self.assertEqual(ws.freeze_panes, "A3")
            self.assertEqual(ws.max_column, 18)
            self.assertEqual(ws["A3"].value, "E001S001C001")
            self.assertEqual(ws["K2"].value, "视频提示词")
            self.assertNotIn("时长", [ws.cell(2, col).value for col in range(1, 19)])
            self.assertIn("50mm", ws["K3"].value)
            self.assertEqual(ws["L3"].value, None)
            self.assertEqual(ws["M3"].value, None)

    def test_v2_requires_beats_three_layer_sound_and_camera_segments(self) -> None:
        self.assertEqual(compiler.validate_records(package(), v2_records()), [])
        value = v2_records(); value["records"][0]["performance_beats"][0]["camera_segment_id"] = "missing"
        errors = compiler.validate_records(package(), value)
        self.assertTrue(any("unknown camera segment" in error for error in errors))

    def test_v2_renders_structured_sound_and_segment_camera_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root / "reference.xlsx"; output = root / "review.xlsx"
            reference(source)
            compiler.render_workbook(package(), v2_records(), source, output)
            ws = load_workbook(output)["S001_测试"]
            self.assertIn("全场环境声：", ws["J3"].value)
            self.assertIn("单点音效：", ws["J3"].value)
            self.assertIn("静默／留白：", ws["J3"].value)
            self.assertEqual(ws["G3"].value, "50mm定焦镜头（全画幅等效）")

    def test_reference_audit_content_never_becomes_current_project_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root / "reference.xlsx"; output = root / "review.xlsx"
            reference(source, title="其他项目的版式")
            template = load_workbook(source)
            for name in ("决定日志", "来源追踪"):
                for address in ("A1", "A2", "A3", "F8"):
                    template[name][address] = "OTHER-PROJECT-PRIVATE-SENTINEL"
            template.save(source)
            compiler.render_workbook(package(), records(), source, output)
            result = load_workbook(output)
            values = [cell.value for sheet in result for row in sheet for cell in row]
            self.assertNotIn("OTHER-PROJECT-PRIVATE-SENTINEL", values)
            self.assertFalse(any(cell.value is not None for row in result["决定日志"].iter_rows(min_row=3) for cell in row))
            trace = list(result["来源追踪"].iter_rows(min_row=3, values_only=True))
            self.assertEqual(len(trace), 3)
            self.assertEqual(trace[0][2:4], ("scn-001", "1镜封存包"))
            self.assertEqual(trace[1][3], "其他项目的版式")
            self.assertIn("已编译 1/1 镜", trace[2][4])

    def test_audit_uses_current_sealed_decisions_scene_and_record_counts(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root / "reference.xlsx"; output = root / "review.xlsx"
            reference(source)
            current = package(); scene = current["scenes"][0]
            scene.update(scene_id="scn-current", production_scene_code="E007S021", worksheet_name="当前场次")
            scene["shots"][0]["shot_id"] = "E007S021C001"
            second = dict(scene["shots"][0], shot_id="E007S021C002")
            scene["shots"].append(second)
            current["decision_log"] = [{
                "decision_id": "current-choice", "topic": "当前议题", "positions": [{"role": "当前岗位"}],
                "chief_recommendation": "当前建议", "tradeoff": "当前取舍", "user_decision": "当前实际决定",
                "impact": "当前影响", "status": "user_decided",
            }]
            current["integrity"]["canonical_json_sha256"] = compiler.package_digest(current)
            payload = v2_records(); payload["records"][0]["shot_id"] = "E007S021C001"
            payload["knowledge"] = {"sources": [{
                "source_id": "current-source", "title": "当前资料", "source_sha256": "source-hash",
                "publication_status": "published", "activation": "inactive",
            }]}
            compiler.render_workbook(current, payload, source, output)
            result = load_workbook(output)
            self.assertTrue(result["当前场次"]["A1"].value.startswith("E007-S021｜"))
            self.assertIsNone(result["当前场次"]["K4"].value)
            decisions = list(result["决定日志"].iter_rows(min_row=3, values_only=True))
            self.assertEqual(len(decisions), 1)
            self.assertEqual(decisions[0][0], "current-choice")
            self.assertEqual(decisions[0][5], "当前实际决定")
            self.assertEqual(json.loads(decisions[0][2]), [{"role": "当前岗位"}])
            trace = list(result["来源追踪"].iter_rows(min_row=3, values_only=True))
            self.assertEqual(trace[0][2:4], ("scn-current", "2镜封存包"))
            self.assertIn("已编译 1/2 镜", trace[2][4])
            self.assertEqual(trace[3][0], "current-source")
            self.assertEqual(trace[3][3], "当前资料｜当前项目未激活")

    def test_v2_multicamera_requires_cut_markers(self) -> None:
        value = v2_records(); record = value["records"][0]
        record["camera_segments"].append({
            "segment_id": "B", "framing": "近景", "movement": "轻推", "angle": "平视",
            "camera_position": "甲肩后看乙", "lens": "65mm定焦镜头（全画幅等效）", "axis": "门床轴线",
        })
        record["performance_beats"].append({
            "beat_id": "B02", "trigger": "乙停下后", "dialogue_fragment": "", "visible_action": "乙抬眼",
            "micro_expression": "眼睑轻颤", "voice": "无台词", "end_state": "两人对视", "camera_segment_id": "B",
        })
        record["prompt"] += "65mm定焦镜头（全画幅等效）拍摄乙抬眼，眼睑轻颤，两人对视。"
        errors = compiler.validate_records(package(), value)
        self.assertTrue(any("镜头切换" in error for error in errors))

    def test_v2_validates_knowledge_activation_records(self) -> None:
        value = v2_records()
        value["knowledge"] = {"sources": [{
            "source_id": "dks-example", "title": "情绪导演", "source_sha256": "abc",
            "publication_status": "published", "activation": "active",
        }]}
        self.assertEqual(compiler.validate_records(package(), value), [])
        value["knowledge"]["sources"][0]["activation"] = "maybe"
        self.assertTrue(any("activation" in error for error in compiler.validate_records(package(), value)))


if __name__ == "__main__":
    unittest.main()
