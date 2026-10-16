#!/usr/bin/env python3
"""Validate prompt records and render a Chinese Seedance shot-list workbook."""

from __future__ import annotations

import argparse
from copy import copy, deepcopy
import hashlib
import json
from pathlib import Path
import re
from typing import Any


PACKAGE_SCHEMA = "SceneDirectionPackageV1"
RECORD_SCHEMA_V1 = "SeedanceShotPromptRecordCollectionV1"
RECORD_SCHEMA_V2 = "SeedanceShotPromptRecordCollectionV2"
PROMPT_PACK_VERSION = "seedance-prompt-pack-v3"
PERFORMANCE_PACK_VERSION = "performance-action-realization-pack-v1"
SKILL_VERSION = "0.8.0"
SCENE_HEADERS = {
    "A": "镜头号", "B": "景别", "C": "运镜", "D": "画面演绎",
    "E": "镜头角度", "F": "机位位置", "G": "焦段", "H": "调度与走位",
    "I": "台词／旁白", "J": "音效／环境声", "K": "视频提示词",
    "N": "轴线", "O": "视线", "P": "人物反应", "Q": "表演目标", "R": "戏剧目的",
}
VISUAL_FORBIDDEN = re.compile(r"台词|旁白|音效|环境声|表演要求|提示词|镜头运动|对白")
PROMPT_FORBIDDEN = re.compile(r"@图片|@Image|@image|图生视频|\bI2V\b|首帧参考|负面提示词", re.IGNORECASE)
PROMPT_BOILERPLATE = re.compile(
    r"16\s*:\s*9古装写实电影画面|台词口型与.{0,12}列同步|不新增对白|"
    r"人物位置、视线方向和房间光线保持连续"
)
TIMECODE = re.compile(r"(?:\[\s*)?\d+(?:\.\d+)?\s*(?:s|秒)\s*[-—–至到]\s*\d+(?:\.\d+)?\s*(?:s|秒)", re.IGNORECASE)


def canonical_bytes(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def hash_json(value: Any) -> str:
    return hashlib.sha256(canonical_bytes(value)).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def package_digest(package: dict[str, Any]) -> str:
    value = deepcopy(package)
    value["integrity"] = {"canonical_json_sha256": None}
    return hash_json(value)


def load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path} must contain a JSON object")
    return value


def all_shots(package: dict[str, Any]) -> list[dict[str, Any]]:
    return [shot for scene in package.get("scenes", []) for shot in scene.get("shots", [])]


def validate_package(package: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if package.get("schema_version") != PACKAGE_SCHEMA:
        errors.append(f"schema_version must be {PACKAGE_SCHEMA}")
    if package.get("status") != "sealed":
        errors.append("only a sealed SceneDirectionPackageV1 may be compiled")
    scenes = package.get("scenes")
    if not isinstance(scenes, list) or not 1 <= len(scenes) <= 2:
        errors.append("package must contain one or two scenes")
        scenes = []
    recorded = package.get("integrity", {}).get("canonical_json_sha256")
    if recorded != package_digest(package):
        errors.append("canonical_json_sha256 does not match sealed package")
    seen: set[str] = set()
    for shot in all_shots(package):
        shot_id = shot.get("shot_id")
        if not isinstance(shot_id, str) or not re.fullmatch(r"[A-Z][A-Z0-9]*S\d{3}C\d{3}", shot_id):
            errors.append(f"invalid shot ID: {shot_id}")
            continue
        if shot_id in seen:
            errors.append(f"duplicate shot ID: {shot_id}")
        seen.add(shot_id)
        for field in ("framing", "movement", "angle", "camera_position", "lens_tendency", "blocking", "axis", "eyeline", "reaction", "performance_target", "dramatic_purpose", "first_state", "last_state"):
            if not str(shot.get(field, "")).strip():
                errors.append(f"{shot_id}: missing {field}")
        if not str(shot.get("dialogue", shot.get("sound", ""))).strip():
            errors.append(f"{shot_id}: missing dialogue source")
        if not str(shot.get("sound_effects", shot.get("sound", ""))).strip():
            errors.append(f"{shot_id}: missing sound source")
    return sorted(set(errors))


def _normalize_text(value: str) -> str:
    return re.sub(r"[\s‘’“”\"'：:，,。！？!?……；;]", "", value)


def _dialogue_content(value: str) -> str:
    value = value.strip()
    if value.startswith("无台词"):
        return ""
    return re.sub(r"^[^：:]{1,20}[：:]", "", value, count=1).strip(" \t\r\n‘’“”\"'")


def _sound_clauses(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[；;]", value) if part.strip()]


def _source_spoken_text(value: str) -> str:
    if value.strip().startswith("无台词"):
        return ""
    quoted = re.findall(r"[‘“]([^’”]+)[’”]", value)
    return "".join(quoted) if quoted else _dialogue_content(value)


def _validate_v2_record(shot_id: str, shot: dict[str, Any], record: dict[str, Any], prompt: str) -> list[str]:
    errors: list[str] = []
    beats = record.get("performance_beats")
    segments = record.get("camera_segments")
    sound = record.get("sound_design")
    if not isinstance(beats, list) or not beats:
        return [f"{shot_id}: V2 performance_beats must be a non-empty list"]
    if not isinstance(segments, list) or not segments:
        return [f"{shot_id}: V2 camera_segments must be a non-empty list"]
    if not isinstance(sound, dict):
        return [f"{shot_id}: V2 sound_design must be an object"]

    segment_ids: set[str] = set()
    for segment in segments:
        segment_id = str(segment.get("segment_id", "")).strip()
        if not segment_id or segment_id in segment_ids:
            errors.append(f"{shot_id}: camera segment IDs must be non-empty and unique")
            continue
        segment_ids.add(segment_id)
        for field in ("framing", "movement", "angle", "camera_position", "lens", "axis"):
            if not str(segment.get(field, "")).strip():
                errors.append(f"{shot_id}/{segment_id}: missing camera {field}")
        lens = _normalize_text(str(segment.get("lens", "")))
        if lens and lens not in _normalize_text(prompt):
            errors.append(f"{shot_id}/{segment_id}: prompt must include the segment lens")

    beat_ids: set[str] = set()
    used_segments: set[str] = set()
    for beat in beats:
        beat_id = str(beat.get("beat_id", "")).strip()
        if not beat_id or beat_id in beat_ids:
            errors.append(f"{shot_id}: performance beat IDs must be non-empty and unique")
            continue
        beat_ids.add(beat_id)
        for field in ("trigger", "visible_action", "micro_expression", "voice", "end_state", "camera_segment_id"):
            if not str(beat.get(field, "")).strip():
                errors.append(f"{shot_id}/{beat_id}: missing performance {field}")
        segment_id = str(beat.get("camera_segment_id", ""))
        if segment_id not in segment_ids:
            errors.append(f"{shot_id}/{beat_id}: unknown camera segment {segment_id}")
        used_segments.add(segment_id)
        fragment = str(beat.get("dialogue_fragment", "")).strip()
        if fragment and _normalize_text(fragment) not in _normalize_text(prompt):
            errors.append(f"{shot_id}/{beat_id}: prompt must embed the dialogue fragment")
    source_spoken = _normalize_text(_source_spoken_text(str(shot.get("dialogue", ""))))
    beat_spoken = _normalize_text("".join(str(beat.get("dialogue_fragment", "")) for beat in beats))
    if source_spoken and source_spoken != beat_spoken:
        errors.append(f"{shot_id}: performance beats must preserve the complete spoken text in order")
    if used_segments != segment_ids:
        errors.append(f"{shot_id}: every camera segment must be bound to at least one performance beat")
    if len(segments) > 1 and prompt.count("镜头切换") < len(segments) - 1:
        errors.append(f"{shot_id}: each additional camera segment requires a visible 镜头切换 marker")

    ambient = str(sound.get("ambient_bed", "")).strip()
    spot = sound.get("spot_effects")
    silence = sound.get("silence_beats")
    if not ambient:
        errors.append(f"{shot_id}: sound_design.ambient_bed is required")
    if not isinstance(spot, list) or not spot:
        errors.append(f"{shot_id}: sound_design.spot_effects must be a non-empty list")
        spot = []
    if not isinstance(silence, list) or not silence:
        errors.append(f"{shot_id}: sound_design.silence_beats must be a non-empty list")
        silence = []
    for item in [ambient, *map(str, spot), *map(str, silence)]:
        if item and _normalize_text(item) not in _normalize_text(prompt):
            errors.append(f"{shot_id}: prompt must embed structured sound: {item}")
    return errors


def validate_records(package: dict[str, Any], payload: dict[str, Any], *, require_all: bool = False) -> list[str]:
    errors = validate_package(package)
    schema = payload.get("schema_version")
    if schema not in {RECORD_SCHEMA_V1, RECORD_SCHEMA_V2}:
        errors.append(f"prompt record schema_version must be {RECORD_SCHEMA_V1} or {RECORD_SCHEMA_V2}")
    if schema == RECORD_SCHEMA_V2 and isinstance(payload.get("knowledge"), dict):
        sources = payload["knowledge"].get("sources", [])
        if not isinstance(sources, list):
            errors.append("knowledge.sources must be a list")
        else:
            source_ids: set[str] = set()
            for source in sources:
                if not isinstance(source, dict):
                    errors.append("knowledge.sources entries must be objects")
                    continue
                source_id = str(source.get("source_id", "")).strip()
                if not source_id or source_id in source_ids:
                    errors.append("knowledge.sources source_id must be non-empty and unique")
                source_ids.add(source_id)
                if source.get("activation") not in {"active", "inactive"}:
                    errors.append(f"{source_id or 'knowledge source'}: activation must be active or inactive")
                for field in ("title", "source_sha256", "publication_status"):
                    if not str(source.get(field, "")).strip():
                        errors.append(f"{source_id or 'knowledge source'}: missing {field}")
    records = payload.get("records")
    if not isinstance(records, list) or not records:
        errors.append("records must be a non-empty list")
        records = []
    shots = {shot["shot_id"]: shot for shot in all_shots(package) if isinstance(shot.get("shot_id"), str)}
    seen: set[str] = set()
    for record in records:
        shot_id = record.get("shot_id")
        if shot_id not in shots:
            errors.append(f"unknown shot ID: {shot_id}")
            continue
        if shot_id in seen:
            errors.append(f"duplicate prompt record: {shot_id}")
        seen.add(shot_id)
        visual = str(record.get("visual", "")).strip()
        prompt = str(record.get("prompt", "")).strip()
        if not visual:
            errors.append(f"{shot_id}: visual is required")
        elif VISUAL_FORBIDDEN.search(visual):
            errors.append(f"{shot_id}: visual contains dialogue, sound, performance, or prompt instructions")
        if not prompt:
            errors.append(f"{shot_id}: prompt is required")
            continue
        if PROMPT_FORBIDDEN.search(prompt):
            errors.append(f"{shot_id}: prompt contains forbidden image-reference or negative-prompt syntax")
        if PROMPT_BOILERPLATE.search(prompt):
            errors.append(f"{shot_id}: prompt contains a generic prefix, continuity placeholder, or audit-style tail")
        if TIMECODE.search(prompt):
            errors.append(f"{shot_id}: prompt contains a hard timecode")
        if not re.search(r"\d+mm|固定机位|镜头", prompt):
            errors.append(f"{shot_id}: prompt must state an executable camera contract")
        normalized_prompt = _normalize_text(prompt)
        if schema == RECORD_SCHEMA_V2:
            errors.extend(_validate_v2_record(shot_id, shots[shot_id], record, prompt))
        else:
            dialogue = _dialogue_content(str(shots[shot_id].get("dialogue", "")))
            if dialogue and _normalize_text(dialogue) not in normalized_prompt:
                errors.append(f"{shot_id}: prompt must embed the complete dialogue at its performance point")
            for clause in _sound_clauses(str(shots[shot_id].get("sound_effects", shots[shot_id].get("sound", "")))):
                if _normalize_text(clause) not in normalized_prompt:
                    errors.append(f"{shot_id}: prompt must embed sound clause: {clause}")
    if require_all and seen != set(shots):
        missing = sorted(set(shots) - seen)
        errors.append("full compilation is missing shot IDs: " + ", ".join(missing))
    return sorted(set(errors))


def visual_only(shot: dict[str, Any]) -> str:
    supplied = str(shot.get("subject_action", "")).strip()
    first_paragraph = supplied.split("\n", 1)[0].strip()
    if not first_paragraph:
        first_paragraph = f"{shot.get('framing', '')}。{shot.get('blocking', '')}。"
    first_paragraph = re.sub(r"动作推进到.*$", "", first_paragraph).strip()
    return f"{first_paragraph}\n画面结束时，{shot.get('last_state', '')}。"


def _copy_cell(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _copy_audit_sheet(source: Any, target: Any) -> None:
    """Reuse only header/body formatting; reference data belongs to another run."""
    from openpyxl.utils import get_column_letter

    for col in range(1, source.max_column + 1):
        letter = get_column_letter(col)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width
    for row in range(1, min(3, source.max_row) + 1):
        target.row_dimensions[row].height = source.row_dimensions[row].height
        for col in range(1, source.max_column + 1):
            src = source.cell(row, col)
            dst = target.cell(row, col)
            _copy_cell(src, dst)
    for merged in source.merged_cells.ranges:
        if merged.max_row == 1:
            target.merge_cells(str(merged))
    target.freeze_panes = "A3"


def _style_scene_sheet(template: Any, target: Any, rows: int) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    for col in range(1, 19):
        letter = get_column_letter(col)
        target.column_dimensions[letter].width = template.column_dimensions[letter].width
    target.column_dimensions["D"].width = 70
    target.column_dimensions["K"].width = 78
    target.column_dimensions["L"].width = 3
    target.column_dimensions["M"].width = 3
    for row in range(1, rows + 3):
        source_row = min(row, template.max_row)
        target.row_dimensions[row].height = template.row_dimensions[source_row].height or (160 if row >= 3 else None)
        if row >= 3:
            target.row_dimensions[row].height = 160
        for col in range(1, 19):
            _copy_cell(template.cell(source_row, col), target.cell(row, col))
            if row >= 2:
                target.cell(row, col).alignment = copy(target.cell(row, col).alignment)
                target.cell(row, col).alignment = Alignment(
                    horizontal=target.cell(row, col).alignment.horizontal or "left",
                    vertical="top" if row >= 3 else "center",
                    wrap_text=True,
                )
    target.merge_cells("A1:K1")
    target.merge_cells("N1:R1")
    target.freeze_panes = "A3"
    target.auto_filter.ref = f"A2:K{rows + 2}"
    target.sheet_view.showGridLines = False


def _segment_cell(record: dict[str, Any] | None, shot: dict[str, Any], field: str) -> str:
    fallback = {"lens": "lens_tendency"}.get(field, field)
    if not record or not isinstance(record.get("camera_segments"), list):
        return str(shot.get(fallback, ""))
    segments = record["camera_segments"]
    if len(segments) == 1:
        return str(segments[0].get(field, shot.get(fallback, "")))
    labels = []
    for index, segment in enumerate(segments):
        label = chr(ord("A") + index) if index < 26 else str(index + 1)
        labels.append(f"{label}段：{segment.get(field, '')}")
    return "\n".join(labels)


def _sound_cell(record: dict[str, Any] | None, shot: dict[str, Any]) -> str:
    if not record or not isinstance(record.get("sound_design"), dict):
        return str(shot.get("sound_effects", shot.get("sound", "")))
    sound = record["sound_design"]
    return (
        f"全场环境声：{sound.get('ambient_bed', '')}\n"
        f"单点音效：{'；'.join(map(str, sound.get('spot_effects', [])))}\n"
        f"静默／留白：{'；'.join(map(str, sound.get('silence_beats', [])))}"
    )


def render_workbook(
    package: dict[str, Any], payload: dict[str, Any], reference: Path, output: Path,
    *, package_path: Path | None = None,
) -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required; load the bundled workspace dependencies") from exc

    errors = validate_records(package, payload, require_all=False)
    if errors:
        raise ValueError("; ".join(errors))
    reference_book = load_workbook(reference)
    template = next((sheet for sheet in reference_book.worksheets
                     if sheet.title not in {"决定日志", "来源追踪"}), None)
    if template is None:
        raise ValueError("reference workbook must contain a scene layout worksheet")
    records = {item["shot_id"]: item for item in payload["records"]}
    shot_count = sum(len(scene["shots"]) for scene in package["scenes"])
    book = Workbook()
    book.remove(book.active)
    for scene in package["scenes"]:
        sheet = book.create_sheet(scene.get("worksheet_name") or scene["production_scene_code"])
        shots = scene["shots"]
        _style_scene_sheet(template, sheet, len(shots))
        review_label = payload.get("review_label", f"已编译 {len(records)}/{shot_count} 镜")
        scene_code = re.sub(r"^(E\d+)(S\d+)$", r"\1-\2", scene["production_scene_code"])
        sheet["A1"] = f"{scene_code}｜Seedance 2.0 文生视频｜16:9｜{review_label}"
        sheet["N1"] = "可参考意见"
        for letter, title in SCENE_HEADERS.items():
            sheet[f"{letter}2"] = title
        for row, shot in enumerate(shots, 3):
            record = records.get(shot["shot_id"])
            values = {
                "A": shot["shot_id"], "B": _segment_cell(record, shot, "framing"), "C": _segment_cell(record, shot, "movement"),
                "D": record["visual"] if record else visual_only(shot),
                "E": _segment_cell(record, shot, "angle"), "F": _segment_cell(record, shot, "camera_position"),
                "G": _segment_cell(record, shot, "lens"), "H": shot.get("blocking"),
                "I": shot.get("dialogue", "无台词。"), "J": _sound_cell(record, shot),
                "K": record["prompt"] if record else "", "N": shot.get("axis"), "O": shot.get("eyeline"),
                "P": shot.get("reaction"), "Q": shot.get("performance_target"), "R": shot.get("dramatic_purpose"),
            }
            for letter, value in values.items():
                sheet[f"{letter}{row}"] = value
            density = max(len(str(values["D"] or "")), len(str(values["K"] or "")), len(str(values["J"] or "")))
            sheet.row_dimensions[row].height = min(420, max(150, 90 + density * 0.32))
    for name in ("决定日志", "来源追踪"):
        target = book.create_sheet(name)
        if name in reference_book.sheetnames:
            _copy_audit_sheet(reference_book[name], target)
    decision = book["决定日志"]
    decision["A1"] = "当前封存导演包的决定日志"
    decision_fields = ("decision_id", "topic", "positions", "chief_recommendation", "tradeoff", "user_decision", "impact", "status")
    decision_headers = ("决定编号", "议题", "岗位立场", "总导演建议", "取舍", "用户决定", "影响", "状态")
    for col, title in enumerate(decision_headers, 1):
        decision.cell(2, col, title)
    decisions = package.get("decision_log", [])
    for row, item in enumerate(decisions, 3):
        for col, key in enumerate(decision_fields, 1):
            if row > 3:
                _copy_cell(decision.cell(3, col), decision.cell(row, col))
            value = item.get(key)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            decision.cell(row, col, value)
    decision.auto_filter.ref = f"A2:H{max(2, len(decisions) + 2)}"
    trace = book["来源追踪"]
    trace["A1"] = "当前编译的来源追踪"
    for col, title in enumerate(("来源", "版本或哈希", "适用范围", "内容", "说明"), 1):
        trace.cell(2, col, title)
    package_hash = sha256_file(package_path) if package_path else package["integrity"]["canonical_json_sha256"]
    scene_ids = "/".join(str(scene.get("scene_id", scene["production_scene_code"])) for scene in package["scenes"])
    entries = [
        ("sealed-scene-direction-package", package_hash, scene_ids, f"{shot_count}镜封存包", f"版本 {package.get('version', '')}；仅派生提示词，不反向修改"),
        ("reference-workbook", sha256_file(reference), "版式", template.title, "仅复用版式，不复制参考项目的决定或来源记录"),
        ("seedance-shot-prompt-compiler", SKILL_VERSION, "提示词编译", "Seedance 2.0 / T2V / 中文 / 16:9", f"通用规则 {PROMPT_PACK_VERSION}；已编译 {len(records)}/{shot_count} 镜"),
    ]
    knowledge = payload.get("knowledge", {}) if isinstance(payload.get("knowledge"), dict) else {}
    for source in knowledge.get("sources", []):
        activation = "当前项目启用" if source.get("activation") == "active" else "当前项目未激活"
        entries.append((
            source.get("source_id", ""), source.get("source_sha256", ""), "导演知识",
            f"{source.get('title', '')}｜{activation}",
            f"发布状态：{source.get('publication_status', '')}",
        ))
    for row, entry in enumerate(entries, 3):
        for col in range(1, 6):
            _copy_cell(trace.cell(max(3, row - 1), col), trace.cell(row, col))
        for col, value in enumerate(entry, 1):
            trace.cell(row, col, value)
        trace.row_dimensions[row].height = 72
    trace.auto_filter.ref = f"A2:E{trace.max_row}"
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    validate = sub.add_parser("validate")
    validate.add_argument("package", type=Path); validate.add_argument("records", type=Path); validate.add_argument("--require-all", action="store_true")
    materialize = sub.add_parser("materialize")
    materialize.add_argument("package", type=Path); materialize.add_argument("records", type=Path)
    materialize.add_argument("--reference", type=Path, required=True); materialize.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    package = load_json(args.package); payload = load_json(args.records)
    errors = validate_records(package, payload, require_all=getattr(args, "require_all", False))
    if errors:
        for error in errors:
            print(f"ERROR: {error}")
        return 1
    if args.command == "materialize":
        render_workbook(package, payload, args.reference, args.output, package_path=args.package)
    print("OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
